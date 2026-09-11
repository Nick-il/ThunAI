import ast
import math
import operator
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy import stats
from langchain_core.tools import tool
import docx  # Added for Word generation

# ============================================================
# NEW: WORD DOCUMENT GENERATOR
# ============================================================
@tool
def generate_word_document(title: str, findings: str, author: str = "ThunAI Agent") -> str:
    """
    Generate a formal Word document (.docx) approval note or report.
    Use this when the user asks to create a report, document, or approval note.
    """
    try:
        # Ensure directory exists
        directory = Path("deliverables")
        directory.mkdir(exist_ok=True)
        
        filename = directory / f"{title.replace(' ', '_')}.docx"
        
        doc = docx.Document()
        doc.add_heading(title, 0)
        
        doc.add_heading('Author', level=1)
        doc.add_paragraph(author)
        
        doc.add_heading('Key Findings & Notes', level=1)
        doc.add_paragraph(findings)
        
        doc.save(filename)
        return f"Success: Word document generated and saved locally at {filename}"
    except Exception as e:
        return f"Document generation error: {e}"

# ============================================================
# SAFE CALCULATOR
# ============================================================
_ALLOWED_OPERATORS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv, ast.Pow: operator.pow, ast.Mod: operator.mod, ast.FloorDiv: operator.floordiv, ast.USub: operator.neg, ast.UAdd: operator.pos}
_ALLOWED_FUNCTIONS = {"sqrt": math.sqrt, "sin": math.sin, "cos": math.cos, "tan": math.tan, "asin": math.asin, "acos": math.acos, "atan": math.atan, "log": math.log, "log10": math.log10, "exp": math.exp, "fabs": math.fabs, "factorial": math.factorial, "ceil": math.ceil, "floor": math.floor, "degrees": math.degrees, "radians": math.radians, "abs": abs}
_ALLOWED_CONSTANTS = {"pi": math.pi, "e": math.e}

def _evaluate(node):
    if isinstance(node, ast.Constant):
        if isinstance(node.value, (int, float)): return node.value
        raise ValueError("Only numerical constants are allowed.")
    if isinstance(node, ast.BinOp):
        operation = type(node.op)
        if operation not in _ALLOWED_OPERATORS: raise ValueError("Operator not allowed.")
        return _ALLOWED_OPERATORS[operation](_evaluate(node.left), _evaluate(node.right))
    if isinstance(node, ast.UnaryOp):
        operation = type(node.op)
        if operation not in _ALLOWED_OPERATORS: raise ValueError("Unary operator not allowed.")
        return _ALLOWED_OPERATORS[operation](_evaluate(node.operand))
    if isinstance(node, ast.Call):
        if not isinstance(node.func, ast.Name): raise ValueError("Invalid function.")
        name = node.func.id
        if name not in _ALLOWED_FUNCTIONS: raise ValueError(f"Function '{name}' not allowed.")
        return _ALLOWED_FUNCTIONS[name](*[_evaluate(arg) for arg in node.args])
    if isinstance(node, ast.Name):
        if node.id in _ALLOWED_CONSTANTS: return _ALLOWED_CONSTANTS[node.id]
        raise ValueError(f"Variable '{node.id}' not allowed.")
    raise ValueError(f"Expression '{type(node).__name__}' not allowed.")

@tool
def calculator(expression: str) -> str:
    """Perform mathematical calculations."""
    try:
        tree = ast.parse(expression, mode="eval")
        result = _evaluate(tree.body)
        return str(int(result) if isinstance(result, float) and result.is_integer() else result)
    except Exception as e: return f"Calculation error: {e}"

# ============================================================
# ALL TOOLS REGISTRY
# ============================================================
tools = [
    generate_word_document, # Our new deliverable tool
    calculator
    # Note: Copy/paste your other existing engineering tools (unit_conversion, reynolds_number, etc.) here just as you had them before!
]